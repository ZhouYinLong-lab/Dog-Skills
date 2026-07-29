"""
Generate a demo PPT showcasing ppt-master's features.
Uses python-pptx to create a presentation about ppt-master itself.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Color Palette ──────────────────────────────────────────────
BG_DARK   = RGBColor(0x0B, 0x0E, 0x14)   # near-black
BG_CARD   = RGBColor(0x15, 0x1A, 0x23)   # card background
ACCENT    = RGBColor(0x3B, 0x82, 0xF6)   # blue accent
ACCENT2   = RGBColor(0x8B, 0x5C, 0xF6)   # purple
ACCENT3   = RGBColor(0x10, 0xB9, 0x81)   # green
ACCENT4   = RGBColor(0xF5, 0x9E, 0x0B)   # amber
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GRAY      = RGBColor(0x9C, 0xA3, 0xAF)
LIGHT     = RGBColor(0xD1, 0xD5, 0xDB)
DARK_TEXT  = RGBColor(0x1F, 0x29, 0x37)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height

# ── Helper Functions ───────────────────────────────────────────

def add_bg(slide, color=BG_DARK):
    """Fill slide background with a solid color."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, left, top, width, height, fill_color=None, border_color=None, radius=None):
    """Add a rounded or plain rectangle."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        left, top, width, height
    )
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    return shape

def add_text_box(slide, left, top, width, height, text, font_size=18, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    """Add a text box with single-style text."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_rich_text(slide, left, top, width, height, runs_list, alignment=PP_ALIGN.LEFT):
    """
    Add text box with multiple styled runs.
    runs_list: list of (text, font_size, color, bold) tuples
    """
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    for i, (text, font_size, color, bold) in enumerate(runs_list):
        if i == 0:
            run = p.runs[0] if p.runs else p.add_run()
        else:
            run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = "Microsoft YaHei"
    return txBox

def add_icon_circle(slide, left, top, size, color, label, sublabel=""):
    """Add a circular icon with label."""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # Label
    add_text_box(slide, left, top + size + Inches(0.15), size, Inches(0.5),
                 label, font_size=13, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    if sublabel:
        add_text_box(slide, left, top + size + Inches(0.55), size, Inches(0.4),
                     sublabel, font_size=10, color=GRAY, alignment=PP_ALIGN.CENTER)

def add_divider(slide, left, top, width, color=ACCENT):
    """Add a thin horizontal divider line."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()

def add_badge(slide, left, top, text, color=ACCENT):
    """Add a small badge/pill."""
    shape = add_rect(slide, left, top, Inches(1.2), Inches(0.35), fill_color=color, radius=True)
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = "Microsoft YaHei"
    p.alignment = PP_ALIGN.CENTER


# ═══════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE
# ═══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide)

# Large accent bar on left
add_rect(slide, Inches(0), Inches(0), Inches(0.12), H, fill_color=ACCENT)

# Badges
add_badge(slide, Inches(1.2), Inches(1.5), "27,000+ ⭐ GitHub")
add_rect(slide, Inches(2.6), Inches(1.5), Inches(1.2), Inches(0.35), fill_color=ACCENT3, radius=True)
tf = add_rect(slide, Inches(2.6), Inches(1.5), Inches(1.2), Inches(0.35), fill_color=ACCENT3, radius=True).text_frame
p = tf.paragraphs[0]; p.text = "MIT 开源"; p.font.size = Pt(9); p.font.color.rgb = WHITE; p.font.bold = True; p.font.name = "Microsoft YaHei"; p.alignment = PP_ALIGN.CENTER

# Title
add_text_box(slide, Inches(1.2), Inches(2.2), Inches(10), Inches(1.2),
             "PPT Master", font_size=64, color=WHITE, bold=True)
add_text_box(slide, Inches(1.2), Inches(3.3), Inches(10), Inches(0.8),
             "AI 生成真正可编辑的 PowerPoint — 原生形状 · 动画 · 语音旁白 · MP4 视频",
             font_size=22, color=GRAY)

add_divider(slide, Inches(1.2), Inches(4.3), Inches(3))

add_text_box(slide, Inches(1.2), Inches(4.7), Inches(8), Inches(0.6),
             "不是截图拼凑的假 PPT  ·  每个元素独立可编辑  ·  数据本地运行零泄露",
             font_size=16, color=LIGHT)

add_text_box(slide, Inches(1.2), Inches(6.2), Inches(6), Inches(0.5),
             "by Hugo He  ·  github.com/hugohe3/ppt-master", font_size=12, color=GRAY)


# ═══════════════════════════════════════════════════════════════
# SLIDE 2 — CORE PIPELINE
# ═══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(10), Inches(0.7),
             "核心流程：从文档到可编辑 PPTX", font_size=36, color=WHITE, bold=True)
add_text_box(slide, Inches(0.8), Inches(1.1), Inches(10), Inches(0.5),
             "SVG 和 DrawingML 本质同源——转换是「方言翻译」而非格式跨越",
             font_size=16, color=GRAY)

# Pipeline steps
steps = [
    ("📄", "输入文档", "PDF / Word\nMarkdown / URL", ACCENT),
    ("🤖", "AI 生成 SVG", "Claude 设计每一页\n布局+配色+图标", ACCENT2),
    ("🐍", "Python 转换", "SVG → DrawingML\n精确工程化转换", ACCENT3),
    ("📊", "输出 .pptx", "原生可编辑\n每个元素独立操作", ACCENT4),
]

for i, (icon, title, desc, color) in enumerate(steps):
    x = Inches(0.8 + i * 3.1)
    y = Inches(2.0)
    # Card
    card = add_rect(slide, x, y, Inches(2.8), Inches(4.5), fill_color=BG_CARD, radius=True)
    # Icon circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.7), y + Inches(0.5), Inches(1.4), Inches(1.4))
    circle.fill.solid(); circle.fill.fore_color.rgb = color; circle.line.fill.background()
    tf = circle.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]; p.text = icon; p.font.size = Pt(36); p.alignment = PP_ALIGN.CENTER
    # Title
    add_text_box(slide, x + Inches(0.2), y + Inches(2.2), Inches(2.4), Inches(0.5),
                 title, font_size=20, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    # Description
    add_text_box(slide, x + Inches(0.2), y + Inches(2.8), Inches(2.4), Inches(1.2),
                 desc, font_size=13, color=GRAY, alignment=PP_ALIGN.CENTER)
    # Arrow between cards (except last)
    if i < 3:
        add_text_box(slide, x + Inches(2.9), y + Inches(1.8), Inches(0.4), Inches(0.5),
                     "→", font_size=28, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

# Bottom note
add_text_box(slide, Inches(0.8), Inches(6.8), Inches(11), Inches(0.5),
             "💡 关键区别：截图流工具把 PNG 塞进 PPT 当背景 → 改不了  |  PPT Master 生成 DrawingML 原生对象 → 随便改",
             font_size=13, color=ACCENT3)


# ═══════════════════════════════════════════════════════════════
# SLIDE 3 — 20 TEMPLATES
# ═══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(10), Inches(0.7),
             "🎨 20 套内置模板", font_size=36, color=WHITE, bold=True)
add_text_box(slide, Inches(0.8), Inches(1.1), Inches(10), Inches(0.5),
             "覆盖商业演示、学术答辩、科技路演全场景，还支持自定义 .pptx 模板",
             font_size=16, color=GRAY)

templates = [
    ("Anthropic 风", "极简科技感", ACCENT),
    ("麦肯锡风", "咨询顾问级", ACCENT2),
    ("Google 风", "Material Design", ACCENT3),
    ("学术答辩风", "严谨规范", ACCENT4),
    ("科技创业风", "Startup Pitch", RGBColor(0xEF, 0x44, 0x44)),
    ("极简风", "Less is More", RGBColor(0x14, 0xB8, 0xA6)),
    ("渐变风", "视觉冲击", RGBColor(0xEC, 0x48, 0x99)),
    ("复古风", "经典韵味", RGBColor(0xD9, 0x77, 0x06)),
    ("更多...", "持续更新中", GRAY),
]

for i, (name, desc, color) in enumerate(templates):
    col = i % 3
    row = i // 3
    x = Inches(0.8 + col * 4.1)
    y = Inches(1.9 + row * 1.75)
    card = add_rect(slide, x, y, Inches(3.8), Inches(1.4), fill_color=BG_CARD, radius=True)
    # Left accent bar
    add_rect(slide, x, y, Inches(0.08), Inches(1.4), fill_color=color)
    add_text_box(slide, x + Inches(0.3), y + Inches(0.2), Inches(3.2), Inches(0.5),
                 name, font_size=20, color=WHITE, bold=True)
    add_text_box(slide, x + Inches(0.3), y + Inches(0.75), Inches(3.2), Inches(0.4),
                 desc, font_size=13, color=GRAY)


# ═══════════════════════════════════════════════════════════════
# SLIDE 4 — MULTI-SIZE + ICONS + CHARTS
# ═══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(10), Inches(0.7),
             "📐 多尺寸输出  ·  🖼️ 6000+ 图标  ·  📊 33 种图表", font_size=36, color=WHITE, bold=True)

# ── LEFT: Sizes ──
add_text_box(slide, Inches(0.8), Inches(1.5), Inches(4), Inches(0.5),
             "支持的输出尺寸", font_size=22, color=ACCENT, bold=True)

sizes = [
    ("16:9", "宽屏演示 · 1920×1080"),
    ("4:3", "传统投影 · 1024×768"),
    ("3:4", "小红书竖版 · 1080×1440"),
    ("1:1", "朋友圈方形 · 1080×1080"),
    ("A4", "打印输出 · 210×297mm"),
    ("自定义", "任意尺寸"),
]
for i, (ratio, desc) in enumerate(sizes):
    y = Inches(2.2 + i * 0.65)
    add_rect(slide, Inches(1.0), y, Inches(3.5), Inches(0.5), fill_color=BG_CARD, radius=True)
    add_text_box(slide, Inches(1.2), y + Inches(0.05), Inches(1.0), Inches(0.4),
                 ratio, font_size=15, color=WHITE, bold=True)
    add_text_box(slide, Inches(2.3), y + Inches(0.05), Inches(2.2), Inches(0.4),
                 desc, font_size=11, color=GRAY)

# ── RIGHT: Icons & Charts ──
add_text_box(slide, Inches(5.5), Inches(1.5), Inches(7), Inches(0.5),
             "内置资源库", font_size=22, color=ACCENT, bold=True)

# Icons card
card1 = add_rect(slide, Inches(5.5), Inches(2.2), Inches(3.5), Inches(2.2), fill_color=BG_CARD, radius=True)
add_text_box(slide, Inches(5.8), Inches(2.4), Inches(3.0), Inches(0.5),
             "🎯 6000+ Tabler 图标", font_size=20, color=WHITE, bold=True)
add_text_box(slide, Inches(5.8), Inches(3.0), Inches(3.0), Inches(1.2),
             "无需下载，AI 按需引用\n箭头·图表·设备·用户·\n社交·文件·编辑·导航\n……覆盖全场景需求",
             font_size=12, color=GRAY)

# Charts card
card2 = add_rect(slide, Inches(9.4), Inches(2.2), Inches(3.5), Inches(2.2), fill_color=BG_CARD, radius=True)
add_text_box(slide, Inches(9.7), Inches(2.4), Inches(3.0), Inches(0.5),
             "📊 33 种图表模板", font_size=20, color=WHITE, bold=True)
add_text_box(slide, Inches(9.7), Inches(3.0), Inches(3.0), Inches(1.2),
             "柱状图·折线图·饼图·雷达图\n甘特图·流程图·组织架构图\n散点图·面积图·环形图\n……静态形状，清晰展示",
             font_size=12, color=GRAY)

# ── BOTTOM NOTE ──
add_text_box(slide, Inches(0.8), Inches(5.8), Inches(11), Inches(0.8),
             "💡 所有资源本地引用，生成过程不联网下载——隐私数据 100% 留在你的电脑上",
             font_size=14, color=ACCENT3)

add_text_box(slide, Inches(0.8), Inches(6.6), Inches(11), Inches(0.5),
             "⚠️ 图表为 PowerPoint 静态形状，不绑定 Excel 数据源（如需动态绑定可手动编辑）",
             font_size=11, color=GRAY)


# ═══════════════════════════════════════════════════════════════
# SLIDE 5 — ANIMATION SYSTEM
# ═══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(10), Inches(0.7),
             "🎬 原生动画系统", font_size=36, color=WHITE, bold=True)
add_text_box(slide, Inches(0.8), Inches(1.1), Inches(10), Inches(0.5),
             "两层动画，DrawingML 原生实现——可在 PowerPoint 中自由编辑每个动画效果",
             font_size=16, color=GRAY)

# ── Layer 1: Transitions ──
add_text_box(slide, Inches(0.8), Inches(1.9), Inches(5), Inches(0.5),
             "① 页面过渡 Slide Transitions", font_size=22, color=ACCENT, bold=True)
add_rect(slide, Inches(0.8), Inches(2.5), Inches(5.8), Inches(2.0), fill_color=BG_CARD, radius=True)
add_text_box(slide, Inches(1.1), Inches(2.7), Inches(5.3), Inches(0.4),
             "✅ 默认开启  ·  Fade 淡入淡出  ·  可自定义切换效果", font_size=13, color=ACCENT3)
add_text_box(slide, Inches(1.1), Inches(3.2), Inches(5.3), Inches(1.0),
             "页面之间的切换动画，让演示更流畅专业。\n默认使用 Fade（淡入淡出）效果，\n也可在 PowerPoint 中手动调整为其他过渡效果。",
             font_size=13, color=GRAY)

# ── Layer 2: Element Animations ──
add_text_box(slide, Inches(0.8), Inches(4.8), Inches(5), Inches(0.5),
             "② 元素入场 Entrance Animations", font_size=22, color=ACCENT2, bold=True)
add_rect(slide, Inches(0.8), Inches(5.4), Inches(5.8), Inches(1.6), fill_color=BG_CARD, radius=True)
add_text_box(slide, Inches(1.1), Inches(5.55), Inches(5.3), Inches(0.4),
             "⚠️ 默认关闭  ·  -a auto 开启  ·  AI 语义分配", font_size=13, color=ACCENT4)
add_text_box(slide, Inches(1.1), Inches(6.0), Inches(5.3), Inches(0.8),
             "逐元素飞入、淡入、缩放等动画。\n设计哲学：默认关闭避免「AI 味」，需要时按需开启。",
             font_size=13, color=GRAY)

# ── Trigger Types ──
add_text_box(slide, Inches(7.5), Inches(1.9), Inches(5), Inches(0.5),
             "三种触发方式", font_size=22, color=ACCENT3, bold=True)

triggers = [
    ("on-click", "点击触发", "适合现场演讲\n演讲者控制节奏", ACCENT),
    ("after-previous", "自动播放", "适合旁白配音\n视频导出场景", ACCENT2),
    ("with-previous", "同步播放", "并行展示\n多元素同时出现", ACCENT3),
]

for i, (code, title, desc, color) in enumerate(triggers):
    y = Inches(2.5 + i * 1.7)
    card = add_rect(slide, Inches(7.5), y, Inches(5.3), Inches(1.4), fill_color=BG_CARD, radius=True)
    add_rect(slide, Inches(7.5), y, Inches(0.08), Inches(1.4), fill_color=color)
    add_badge(slide, Inches(7.9), y + Inches(0.2), code, color)
    add_text_box(slide, Inches(9.3), y + Inches(0.15), Inches(3.3), Inches(0.4),
                 title, font_size=18, color=WHITE, bold=True)
    add_text_box(slide, Inches(7.9), y + Inches(0.65), Inches(4.7), Inches(0.6),
                 desc, font_size=11, color=GRAY)


# ═══════════════════════════════════════════════════════════════
# SLIDE 6 — VOICE NARRATION + VIDEO EXPORT
# ═══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(10), Inches(0.7),
             "🎙️ 语音旁白 + MP4 视频导出", font_size=36, color=WHITE, bold=True)

# Flow
flow_steps = [
    ("📝", "演讲者备注", "AI 为每页生成\n演讲要点", ACCENT),
    ("🔊", "AI 语音合成", "90+ 语言\n自然语音输出", ACCENT2),
    ("📦", "嵌入 PPTX", "音频嵌入文件\n无需外部依赖", ACCENT3),
    ("🎥", "导出 MP4", "PowerPoint 一键\n导出视频文件", ACCENT4),
]

for i, (icon, title, desc, color) in enumerate(flow_steps):
    x = Inches(0.8 + i * 3.1)
    y = Inches(1.6)
    card = add_rect(slide, x, y, Inches(2.8), Inches(3.2), fill_color=BG_CARD, radius=True)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.7), y + Inches(0.4), Inches(1.4), Inches(1.4))
    circle.fill.solid(); circle.fill.fore_color.rgb = color; circle.line.fill.background()
    tf = circle.text_frame; p = tf.paragraphs[0]; p.text = icon; p.font.size = Pt(36); p.alignment = PP_ALIGN.CENTER
    add_text_box(slide, x + Inches(0.2), y + Inches(2.1), Inches(2.4), Inches(0.5),
                 title, font_size=18, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), y + Inches(2.6), Inches(2.4), Inches(0.6),
                 desc, font_size=12, color=GRAY, alignment=PP_ALIGN.CENTER)
    if i < 3:
        add_text_box(slide, x + Inches(2.9), y + Inches(1.3), Inches(0.4), Inches(0.5),
                     "→", font_size=28, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

# Key insight
add_rect(slide, Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.8), fill_color=BG_CARD, radius=True)
add_text_box(slide, Inches(1.2), Inches(5.4), Inches(11), Inches(0.5),
             "💡 动画 + 旁白的完美配合", font_size=22, color=ACCENT4, bold=True)
add_text_box(slide, Inches(1.2), Inches(6.0), Inches(11), Inches(0.8),
             "导出视频时，元素动画自动切换到 after-previous 模式（无需点击触发），\n配合 AI 语音旁白的节奏逐元素呈现——生成可直接上传到 B站/YouTube 的演示视频。",
             font_size=14, color=LIGHT)


# ═══════════════════════════════════════════════════════════════
# SLIDE 7 — COMPARISON
# ═══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(10), Inches(0.7),
             "⚔️ PPT Master vs 截图流工具", font_size=36, color=WHITE, bold=True)

# Table-like layout
headers = ["对比维度", "❌ 截图流工具", "✅ PPT Master"]
col_widths = [Inches(3), Inches(4.5), Inches(4.5)]
col_starts = [Inches(1.0), Inches(4.0), Inches(8.5)]

# Header row
for i, (header, w, x) in enumerate(zip(headers, col_widths, col_starts)):
    color = GRAY if i == 0 else (RGBColor(0xEF, 0x44, 0x44) if i == 1 else ACCENT3)
    add_rect(slide, x, Inches(1.5), w, Inches(0.55), fill_color=BG_CARD, radius=False)
    add_text_box(slide, x + Inches(0.2), Inches(1.55), w - Inches(0.4), Inches(0.45),
                 header, font_size=16, color=color, bold=True)

# Data rows
rows = [
    ("可编辑性", "全是图片背景，改不了", "每个元素独立可编辑 ✅"),
    ("文本修改", "需要删图重做", "直接点击修改文字 ✅"),
    ("图表数据", "无法修改", "形状可调，颜色可改 ✅"),
    ("动画效果", "没有或固定", "原生动画，自由编辑 ✅"),
    ("文件大小", "巨大（全是 PNG）", "小（矢量图形） ✅"),
    ("再编辑成本", "≈ 重做", "≈ 改几个字 ✅"),
    ("数据安全", "上传云端", "100% 本地运行 ✅"),
]

for i, (dim, bad, good) in enumerate(rows):
    y = Inches(2.2 + i * 0.6)
    bg = BG_CARD if i % 2 == 0 else BG_DARK
    for j, (text, x, w) in enumerate(zip([dim, bad, good], col_starts, col_widths)):
        add_rect(slide, x, y, w, Inches(0.5), fill_color=bg, radius=False)
        color = GRAY if j == 0 else (RGBColor(0xEF, 0x44, 0x44) if j == 1 else ACCENT3)
        add_text_box(slide, x + Inches(0.2), y + Inches(0.08), w - Inches(0.4), Inches(0.35),
                     text, font_size=12, color=color)

# Bottom note
add_text_box(slide, Inches(0.8), Inches(6.6), Inches(11), Inches(0.5),
             "结论：截图流 = 一次性的「样子货」  |  PPT Master = 持续可迭代的专业演示资产",
             font_size=16, color=WHITE, bold=True)


# ═══════════════════════════════════════════════════════════════
# SLIDE 8 — INSTALL & USAGE
# ═══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(10), Inches(0.7),
             "🚀 安装与使用", font_size=36, color=WHITE, bold=True)

# ── Install ──
add_text_box(slide, Inches(0.8), Inches(1.4), Inches(5), Inches(0.5),
             "安装方式", font_size=22, color=ACCENT, bold=True)

install_methods = [
    ("Claude Code 插件", "/plugin marketplace add hugohe3/ppt-master\n/plugin install ppt-master@ppt-master", ACCENT),
    ("跨 Agent CLI", "npx skills add hugohe3/ppt-master", ACCENT2),
    ("Python 依赖", "pip install -r requirements.txt\n# Python 3.10+", ACCENT3),
]

for i, (title, cmd, color) in enumerate(install_methods):
    y = Inches(2.1 + i * 1.5)
    card = add_rect(slide, Inches(0.8), y, Inches(5.5), Inches(1.2), fill_color=BG_CARD, radius=True)
    add_rect(slide, Inches(0.8), y, Inches(0.08), Inches(1.2), fill_color=color)
    add_badge(slide, Inches(1.2), y + Inches(0.1), title, color)
    add_text_box(slide, Inches(1.2), y + Inches(0.55), Inches(5.0), Inches(0.6),
                 cmd, font_size=11, color=LIGHT)

# ── Usage ──
add_text_box(slide, Inches(7.2), Inches(1.4), Inches(5), Inches(0.5),
             "常见用法", font_size=22, color=ACCENT, bold=True)

usages = [
    "帮我把这份 PDF 做成 PPT\n→ 麦肯锡风，10 页以内",
    "把这篇 Markdown 转成\n→ 学术答辩风格的幻灯片",
    "用这个 Word 生成投资人 Pitch\n→ 16:9，带动画",
    "给这份 PPT 加上演讲旁白\n→ 导出 MP4 视频",
]

for i, usage in enumerate(usages):
    y = Inches(2.1 + i * 1.15)
    add_rect(slide, Inches(7.2), y, Inches(5.5), Inches(0.95), fill_color=BG_CARD, radius=True)
    add_text_box(slide, Inches(7.5), y + Inches(0.1), Inches(5.0), Inches(0.75),
                 usage, font_size=12, color=LIGHT)

# ── Performance ──
add_text_box(slide, Inches(0.8), Inches(6.0), Inches(11), Inches(0.5),
             "⚡ 性能参考：5 页 ~5-10 分钟  |  10 页 ~10-20 分钟  |  20+ 页建议分批生成  |  推荐 Claude Opus",
             font_size=13, color=GRAY)


# ═══════════════════════════════════════════════════════════════
# SLIDE 9 — BEST PAIRINGS
# ═══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(10), Inches(0.7),
             "🤝 最佳搭配 —— Dog-Skills 生态系统", font_size=36, color=WHITE, bold=True)

pairings = [
    ("humanize-ppt", "演讲结构设计", "先用 humanize-ppt 做 AST 大纲 + 演讲体检，\n确保叙事逻辑站得住", ACCENT),
    ("ppt-master", "渲染输出", "再用 ppt-master 渲染成可编辑 .pptx，\n加动画、旁白、导出视频", ACCENT2),
    ("baoyu-skills", "配图生成", "需要自定义插图时，用 baoyu-image-gen\n生成高质量图片嵌入 PPT", ACCENT3),
    ("markitdown", "文档预处理", "先把 PDF/Word 转成干净 Markdown，\n喂给 ppt-master 效果更好", ACCENT4),
    ("theme-factory", "品牌配色", "用主题工厂生成品牌配色方案，\n作为自定义模板的色彩参考", RGBColor(0xEC, 0x48, 0x99)),
]

for i, (name, role, desc, color) in enumerate(pairings):
    y = Inches(1.5 + i * 1.1)
    card = add_rect(slide, Inches(0.8), y, Inches(11.7), Inches(0.9), fill_color=BG_CARD, radius=True)
    add_rect(slide, Inches(0.8), y, Inches(0.08), Inches(0.9), fill_color=color)
    add_text_box(slide, Inches(1.2), y + Inches(0.1), Inches(2.5), Inches(0.35),
                 name, font_size=16, color=color, bold=True)
    add_badge(slide, Inches(3.8), y + Inches(0.12), role, color)
    add_text_box(slide, Inches(5.2), y + Inches(0.08), Inches(7.0), Inches(0.7),
                 desc, font_size=12, color=GRAY)

# Flow diagram
add_text_box(slide, Inches(0.8), Inches(7.0), Inches(11), Inches(0.4),
             "推荐流水线:  markitdown → humanize-ppt → ppt-master → baoyu-skills（配图） → MP4 发布",
             font_size=13, color=ACCENT, bold=True)


# ═══════════════════════════════════════════════════════════════
# SLIDE 10 — OUTRO
# ═══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

add_rect(slide, Inches(0), Inches(0), Inches(0.12), H, fill_color=ACCENT)

# GitHub stats
add_text_box(slide, Inches(1.5), Inches(1.8), Inches(10), Inches(0.8),
             "PPT Master", font_size=56, color=WHITE, bold=True)

stats = [
    ("⭐ 27,000+", "GitHub Stars"),
    ("🔓 MIT", "开源协议"),
    ("🐍 Python 3.10+", "环境要求"),
    ("💻 本地运行", "零数据泄露"),
]
for i, (value, label) in enumerate(stats):
    x = Inches(1.5 + i * 2.8)
    add_text_box(slide, x, Inches(2.8), Inches(2.5), Inches(0.5),
                 value, font_size=24, color=ACCENT, bold=True)
    add_text_box(slide, x, Inches(3.3), Inches(2.5), Inches(0.4),
                 label, font_size=13, color=GRAY)

add_divider(slide, Inches(1.5), Inches(3.9), Inches(4))

add_text_box(slide, Inches(1.5), Inches(4.3), Inches(10), Inches(0.9),
             '"不是截图拼凑的假 PPT，而是真正可编辑的演示资产。"',
             font_size=22, color=LIGHT)

add_text_box(slide, Inches(1.5), Inches(5.5), Inches(10), Inches(0.6),
             "github.com/hugohe3/ppt-master", font_size=16, color=ACCENT)
add_text_box(slide, Inches(1.5), Inches(6.1), Inches(10), Inches(0.5),
             "本演示稿由 python-pptx 生成  ·  演示 PPT Master v1.0 各项功能  ·  Dog-Skills 收录",
             font_size=11, color=GRAY)


# ═══════════════════════════════════════════════════════════════
# SAVE
# ═══════════════════════════════════════════════════════════════
output_dir = os.path.dirname(os.path.abspath(__file__))
output_path = os.path.join(output_dir, "ppt-master-demo.pptx")
prs.save(output_path)
print(f"✅ Demo PPT saved to: {output_path}")
print(f"   Slides: {len(prs.slides)}")
print(f"   Size: {os.path.getsize(output_path) / 1024:.1f} KB")
